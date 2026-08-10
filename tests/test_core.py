import json
import socket
import threading
import time
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime

from openpyxl import load_workbook

import simtrail.bridge_client as bridge_client
from simtrail.database import RunRepository
from simtrail.exporters import export_comparison_xlsx, export_runs_xlsx
from simtrail.importers import discover_snapshots, load_snapshot
from simtrail.live_connection import WorkbenchConnectionManager
from simtrail.models import Run, diff_runs
from simtrail.project_io import load_project_bundle, save_project_bundle
from simtrail.resources import resource_path
from simtrail.workbench import WorkbenchSystem, load_workbench_scan, suggest_reference, system_to_run


def make_run(name="Baseline", mesh="2 mm", stress="200 MPa"):
    return Run(
        name=name, project="Bracket", analyst="Test", created_at=datetime(2026, 7, 16, 10, 0),
        inputs={"Material": "Steel", "Mesh": mesh},
        results={"Max equivalent stress": stress, "Minimum safety factor": 2.1},
    )


def test_bundled_example_resource_is_available():
    assert resource_path("examples", "workbench-project.simtrail-workbench.json").is_file()


def test_repository_round_trip(tmp_path):
    repository = RunRepository(tmp_path / "runs.db")
    run_id = repository.add(make_run())
    loaded = repository.get(run_id)
    assert loaded is not None
    assert loaded.name == "Baseline"
    assert loaded.inputs["Mesh"] == "2 mm"
    assert repository.count() == 1
    repository.delete(run_id)
    assert repository.count() == 0


def test_diff_reports_only_changes():
    before = make_run()
    after = make_run("Refined", "1 mm", "215 MPa")
    changes = diff_runs(before, after)
    assert [(change.section, change.field) for change in changes] == [
        ("Inputs", "Mesh"), ("Results", "Max equivalent stress")
    ]


def test_snapshot_import_and_discovery(tmp_path):
    snapshot = tmp_path / "nested" / "run.simtrail.json"
    snapshot.parent.mkdir()
    snapshot.write_text(
        '{"name":"R1","project":"P1","inputs":{"Mesh":"2 mm"},'
        '"results":{"Stress":"100 MPa"}}', encoding="utf-8"
    )
    assert discover_snapshots(tmp_path) == [snapshot]
    run = load_snapshot(snapshot)
    assert run.name == "R1"
    assert run.source_path == str(snapshot)


def test_excel_exports(tmp_path):
    before = make_run()
    after = make_run("Refined", "1 mm", "215 MPa")
    register = tmp_path / "register.xlsx"
    comparison = tmp_path / "comparison.xlsx"
    export_runs_xlsx([before, after], register)
    export_comparison_xlsx(before, after, comparison)
    assert load_workbook(register).active.max_row == 3
    compare_sheet = load_workbook(comparison).active
    assert compare_sheet["A1"].value == "SIMTRAIL RUN COMPARISON"
    assert compare_sheet.max_row == 4


def test_workbench_scan_reference_and_conversion():
    project = load_workbench_scan("examples/workbench-project.simtrail-workbench.json")
    assert len(project.systems) == 10
    reference, reason = suggest_reference(project.systems)
    assert reference is not None
    assert reference.system_id == "F"
    assert "derive" in reason
    reference.notes = "Client-requested nominal case"
    run = system_to_run(
        project, reference, "Analyst", "F",
        {"purpose": "Nominal model", "expected_changes": "None",
         "conclusion": "Meets the target", "decision": "Accepted"},
    )
    assert run.status == "Approved reference"
    assert run.inputs["Workbench system"] == "F"
    assert run.inputs["Analyst purpose"] == "Nominal model"
    assert run.inputs["Analyst expected changes"] == "None"
    assert run.inputs["Analyst conclusion"] == "Meets the target"
    assert run.inputs["Analyst Workbench note"] == "Client-requested nominal case"
    assert run.results["Maximum joint force"] == "18.4 kN"


def test_lsdyna_system_is_classified_from_solution_cells():
    system = WorkbenchSystem(
        system_id="SYS 1", name="bullet&steel", analysis_type="LSDYNA",
        state="Up-to-date", state_source="Solution",
        cell_states={"Geometry": "Up-to-date", "Model": "Up-to-date",
                     "Setup": "Up-to-date", "Solution": "Up-to-date"},
    )
    assert system.is_analysis
    geometry = WorkbenchSystem(
        system_id="SYS", name="Geometry", analysis_type="Geometry",
        cell_states={"Geometry": "Up-to-date"},
    )
    assert not geometry.is_analysis


def test_live_workbench_bridge_protocol(tmp_path, monkeypatch):
    payload = json.loads(open("examples/workbench-project.simtrail-workbench.json", encoding="utf-8").read())
    server = socket.socket()
    server.bind(("127.0.0.1", 0))
    server.listen(1)
    port = server.getsockname()[1]

    def respond():
        connection, _ = server.accept()
        with connection:
            assert connection.recv(256).strip() == b"SIMTRAIL_SCAN_V1"
            connection.sendall(json.dumps(payload).encode("utf-8"))
        server.close()

    thread = threading.Thread(target=respond)
    thread.start()
    monkeypatch.setattr(bridge_client, "PORT", port)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    project = bridge_client.scan_connected_workbench(timeout=2)
    thread.join(timeout=2)
    assert project.name == "Rigid Dynamics Variant Study"
    assert len(project.systems) == 10


def test_repository_merges_exact_snapshot_and_tracks_changes(tmp_path):
    repository = RunRepository(tmp_path / "merge.db")
    assert not repository.is_initialized()
    repository.mark_initialized()
    first = make_run()
    _, created = repository.add_or_update(first)
    assert created
    renamed = make_run(name="Baseline renamed")
    renamed.notes = "Reviewed"
    renamed.inputs["Analyst purpose"] = "Confirm the nominal design"
    run_id, created = repository.add_or_update(renamed)
    assert not created
    assert repository.count() == 1
    assert repository.get(run_id).name == "Baseline renamed"
    assert repository.get(run_id).inputs["Analyst purpose"] == "Confirm the nominal design"
    _, created = repository.add_or_update(make_run(stress="205 MPa"))
    assert created
    assert repository.count() == 2
    assert repository.clear() == 2
    assert repository.count() == 0
    assert repository.is_initialized()


def test_save_and_reopen_project_bundle(tmp_path):
    path = tmp_path / "bracket.simtrail-project.json"
    save_project_bundle([make_run()], path, "Bracket study")
    title, runs = load_project_bundle(path)
    assert title == "Bracket study"
    assert len(runs) == 1
    assert runs[0].inputs["Material"] == "Steel"


def test_persistent_workbench_protocol_round_trip():
    from PySide6.QtCore import QCoreApplication
    from PySide6.QtNetwork import QLocalSocket

    app = QCoreApplication.instance() or QCoreApplication([])
    manager = WorkbenchConnectionManager("SimTrail.Test." + uuid.uuid4().hex)
    received = []
    completed = []
    manager.project_received.connect(received.append)
    manager.command_completed.connect(completed.append)
    assert manager.start()
    client = QLocalSocket()
    client.connectToServer(manager.pipe_name)
    assert client.waitForConnected(1000)
    hello = {"type": "hello", "protocol": 1, "instance_id": "wb-1", "process_id": 123,
             "connector_version": 2, "capabilities": ["scan_project", "extract_results"],
             "ansys_version": "2026 R1", "project_name": "Bracket", "project_path": "P.wbpj"}
    client.write((json.dumps(hello) + "\n").encode())
    client.flush()

    deadline = time.monotonic() + 1
    while not manager.connected and time.monotonic() < deadline:
        app.processEvents()
    assert manager.active_instance.project_name == "Bracket"
    assert manager.active_instance.connector_version == 2
    assert manager.active_instance.supports("extract_results")

    request_id = manager.request_scan()
    deadline = time.monotonic() + 1
    incoming = b""
    while b'"command":"scan_project"' not in incoming and time.monotonic() < deadline:
        app.processEvents()
        client.waitForReadyRead(20)
        incoming += bytes(client.readAll())
    command_line = next(line for line in incoming.splitlines() if b'"command":"scan_project"' in line)
    command = json.loads(command_line)
    assert command["request_id"] == request_id
    payload = json.loads(open("examples/workbench-project.simtrail-workbench.json", encoding="utf-8").read())
    response = {"type": "response", "request_id": request_id, "ok": True, "result": payload}
    client.write((json.dumps(response) + "\n").encode())
    client.flush()
    deadline = time.monotonic() + 1
    while not received and time.monotonic() < deadline:
        app.processEvents()
    assert received[0].name == "Rigid Dynamics Variant Study"
    extraction_id = manager.extract_results(["F", "G"])
    deadline = time.monotonic() + 1
    extraction_incoming = b""
    while b'"command":"extract_results"' not in extraction_incoming and time.monotonic() < deadline:
        app.processEvents()
        client.waitForReadyRead(20)
        extraction_incoming += bytes(client.readAll())
    extraction_line = next(
        line for line in extraction_incoming.splitlines() if b'"command":"extract_results"' in line
    )
    extraction_command = json.loads(extraction_line)
    assert extraction_command["payload"]["system_ids"] == ["F", "G"]
    client.write((json.dumps({
        "type": "response", "request_id": extraction_id, "ok": True,
        "result": {"systems": [{"id": "F", "results": {"Joint force · Maximum": "18 kN"}}]},
    }) + "\n").encode())
    client.flush()
    deadline = time.monotonic() + 1
    while not any(item.get("request_id") == extraction_id for item in completed) and time.monotonic() < deadline:
        app.processEvents()
    assert any(item.get("command") == "extract_results" for item in completed)
    client.disconnectFromServer()
    manager.stop()


def test_act_extension_manifest_has_project_lifecycle():
    root = ET.parse("ansys_extension/SimTrailConnector.xml").getroot()
    interface = root.find("interface")
    assert interface.attrib["context"] == "Project"
    callbacks = interface.find("callbacks")
    assert callbacks.find("oninit").text == "on_init"
    assert callbacks.find("onterminate").text == "on_terminate"
