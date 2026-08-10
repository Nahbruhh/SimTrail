from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Run, diff_runs

NAVY = "17243A"
TEAL = "20B8A6"
PALE = "E8F8F5"


def export_runs_xlsx(runs: list[Run], path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Run register"
    input_keys = sorted({key for run in runs for key in run.inputs})
    result_keys = sorted({key for run in runs for key in run.results})
    headers = ["Run", "Project", "Analyst", "Captured", "Status", "Solver"] + input_keys + result_keys
    sheet.append(headers)
    for run in runs:
        sheet.append([run.name, run.project, run.analyst, run.created_at.strftime("%Y-%m-%d %H:%M"),
                      run.status, run.solver] + [run.inputs.get(k, "") for k in input_keys]
                     + [run.results.get(k, "") for k in result_keys])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    for idx, column in enumerate(sheet.columns, 1):
        width = min(36, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[get_column_letter(idx)].width = width
    workbook.save(path)


def export_comparison_xlsx(before: Run, after: Run, path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Comparison"
    sheet.append(["SIMTRAIL RUN COMPARISON", "", ""])
    sheet.append(["Field", before.name, after.name])
    for difference in diff_runs(before, after):
        sheet.append([f"{difference.section} / {difference.field}", difference.before, difference.after])
    sheet.merge_cells("A1:C1")
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    sheet["A1"].alignment = Alignment(vertical="center")
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A3"
    for index, width in enumerate((38, 32, 32), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(path)

